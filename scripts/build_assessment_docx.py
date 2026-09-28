"""Build a lab-assessment docx in the exact look of the reference file:
letter page, 0.65 in side and 0.55 in top and bottom margins, Times New
Roman, centred bold 11 pt header block, 9.5 pt body at single spacing with
1 pt after, bold 9.5 pt section headings with 2 pt before, bullets as text
lines inside one paragraph."""
import re
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

src = Path(sys.argv[1])
out = Path(sys.argv[2])
figdir = src.parent

BODY_PT = 9.5
HEAD_PT = 11

d = docx.Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.65)
sec.top_margin = sec.bottom_margin = Inches(0.55)

normal = d.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(BODY_PT)
# east-asian font slot too, or Word falls back to Calibri for some runs
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
if rfonts is None:
    from docx.oxml import OxmlElement
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + attr.split(":")[1],
               "Times New Roman")
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(1)
normal.paragraph_format.line_spacing = 1.0


def add_runs(par, text, size=BODY_PT, bold_all=False):
    """Text with **bold** spans."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_all or (part.startswith("**") and part.endswith("**"))
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
        r = par.add_run(part)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = "Times New Roman"


def para(text="", size=BODY_PT, bold=False, align=None, before=0, after=1, indent=None):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
    if indent is not None:
        pf.left_indent = Inches(indent)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        add_runs(p, text, size=size, bold_all=bold)
    return p


def add_table(rows):
    """A pipe table, ruled and full width, at one point under body size."""
    tb = d.add_table(rows=len(rows), cols=len(rows[0]))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = tb._tbl.tblPr
    w = OxmlElement("w:tblW"); w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
    tbl_pr.append(w)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080")
        borders.append(e)
    tbl_pr.append(borders)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = tb.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text, size=BODY_PT - 1, bold_all=(i == 0))
            if i == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EDEDED")
                tc_pr.append(shd)
    para("", after=3)


lines = src.read_text(encoding="utf-8").splitlines()

# ---- header block: the leading lines up to the first bold title line
i = 0
header = []
while i < len(lines):
    ln = lines[i].strip()
    if ln.startswith("**"):
        break
    if ln:
        header.append(ln)
    i += 1
for h in header:
    para(h, size=HEAD_PT, bold=True, align="center", after=0)
para("", after=0)

# project title: the first bold line, centred
title = lines[i].strip().strip("*")
para(title, size=BODY_PT + 1, bold=True, align="center", before=4, after=6)
i += 1

# ---- body
bullets = []
numbers = []
table_rows = []


def flush_lists():
    global bullets, numbers, table_rows
    if table_rows:
        add_table(table_rows)
        table_rows = []
    if bullets:
        p = para(indent=0.35, after=1)
        for k, b in enumerate(bullets):
            if k:
                p.add_run().add_break(WD_BREAK.LINE)
            add_runs(p, "• " + b)
        bullets = []
    if numbers:
        p = para(indent=0.2, after=1)
        for k, n in enumerate(numbers):
            if k:
                p.add_run().add_break(WD_BREAK.LINE)
            add_runs(p, n)
        numbers = []


while i < len(lines):
    ln = lines[i].rstrip()
    s = ln.strip()
    i += 1
    if not s:
        flush_lists()
        continue
    if s.startswith("|") and s.endswith("|"):
        cells = [c.strip() for c in s.strip("|").split("|")]
        if not all(set(c) <= set("-: ") for c in cells):
            table_rows.append(cells)
        continue
    m = re.match(r"!\[\]\(([^)]+)\)", s)
    if m:
        flush_lists()
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(figdir / m.group(1)), width=Inches(6.6))
        continue
    if s.startswith("- "):
        numbers and flush_lists()
        bullets.append(s[2:].strip())
        continue
    if re.match(r"^\d+\.\s", s) and not s.startswith("**"):
        bullets and flush_lists()
        numbers.append(s)
        continue
    flush_lists()
    if s.startswith("**") and s.endswith("**") and s.count("**") == 2:
        para(s[2:-2], bold=True, before=2, after=1)
    elif s.startswith("Figure ") and "." in s[:12]:
        para(s, align="center", after=3)
    else:
        para(s)

flush_lists()
d.save(str(out))
print("wrote", out, "|", len(d.paragraphs), "paragraphs")
